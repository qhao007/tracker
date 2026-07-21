"""
Validation Summary PPT 生成器单元测试 - v0.16.0

直接测试 app/ppt_report.py 的 generate_validation_summary_ppt()：
- KPI 字段填充（简单占比口径）
- 环比行（有/无快照数据）
- 评语规则生成
- 图表占位图替换

运行命令:
    cd dev && PYTHONPATH=. pytest tests/test_api/test_ppt_report.py -v
"""

import base64
import io
import sys
import os

import pytest

# 确保导入路径正确
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from app.ppt_report import generate_validation_summary_ppt
from pptx import Presentation

# 1x1 透明 PNG
SAMPLE_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)

# 对齐样张数值的合成数据
BASE_REPORT = {
    'project': {'name': 'EX1', 'start_date': '2026-01-01', 'end_date': '2026-12-31'},
    'overview': {'total_cp': 254, 'covered_cp': 236, 'coverage_rate': 92.9},
    'by_priority': {
        'P0': {'total': 78, 'covered': 78, 'rate': 100.0},
        'P1': {'total': 176, 'covered': 158, 'rate': 89.8},
        'P2': {'total': 0, 'covered': 0, 'rate': 0},
    },
    'coverage_matrix': {'matrix': {
        'APCS': {'P0': {'covered': 2, 'total': 2}, 'P1': {'covered': 1, 'total': 3}},
        'PMA': {'P0': {'covered': 1, 'total': 1}, 'P1': {'covered': 5, 'total': 8}},
        'CCB': {'P0': {'covered': 10, 'total': 10}, 'P1': {'covered': 40, 'total': 41}},
    }},
}


def _walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # GROUP
            yield from _walk(sh.shapes)


def _open(ppt_bytes):
    return Presentation(io.BytesIO(ppt_bytes))


def _all_text(prs):
    texts = []
    for sh in _walk(prs.slides[0].shapes):
        if sh.has_text_frame:
            texts.append(sh.text_frame.text)
    return "\n".join(texts)


def _pictures(prs):
    return [sh for sh in _walk(prs.slides[0].shapes) if sh.shape_type == 13]


class TestGenerateValidationSummaryPPT:
    """PPT 生成功能"""

    def test_basic_fields(self):
        """PPT-UNIT-001: 标题与 KPI 数值正确填充"""
        prs = _open(generate_validation_summary_ppt(BASE_REPORT))
        text = _all_text(prs)

        assert 'EX1-Functional Validation Summary' in text
        assert '92.9%' in text
        assert '236' in text
        assert '18' in text  # 未覆盖 = 254 - 236
        assert '78 / 176' in text
        assert 'P0=30.7%，P1=69.3%' in text

    def test_delta_line_present(self):
        """PPT-UNIT-002: 有快照环比时显示 较 M/D +x.xpp"""
        prs = _open(generate_validation_summary_ppt(
            BASE_REPORT, coverage_delta={'prev_date': '2026-05-28', 'delta_pp': 0.8}))
        assert '较 5/28 +0.8pp' in _all_text(prs)

    def test_delta_line_negative(self):
        """PPT-UNIT-003: 环比下降时显示负号"""
        prs = _open(generate_validation_summary_ppt(
            BASE_REPORT, coverage_delta={'prev_date': '2026-05-28', 'delta_pp': -1.2}))
        assert '较 5/28 -1.2pp' in _all_text(prs)

    def test_delta_removed_when_none(self):
        """PPT-UNIT-004: 无环比数据时删除环比行"""
        prs = _open(generate_validation_summary_ppt(BASE_REPORT, coverage_delta=None))
        assert 'pp' not in _all_text(prs)

    def test_commentary_low_p1(self):
        """PPT-UNIT-005: 低 P1 覆盖 Feature 点名 + P0 全部覆盖"""
        prs = _open(generate_validation_summary_ppt(BASE_REPORT))
        assert 'APCS、PMA 的 P1 功能覆盖率不高；但 P0 核心功能已全部覆盖' in _all_text(prs)

    def test_commentary_p0_incomplete(self):
        """PPT-UNIT-006: P0 未全部覆盖时提示补齐"""
        report = dict(BASE_REPORT)
        report['coverage_matrix'] = {'matrix': {
            'DDR': {'P0': {'covered': 1, 'total': 2}, 'P1': {'covered': 9, 'total': 10}},
        }}
        prs = _open(generate_validation_summary_ppt(report))
        assert 'P0 核心功能尚未全部覆盖，需优先补齐' in _all_text(prs)

    def test_commentary_all_normal(self):
        """PPT-UNIT-007: 无 P0 且无低覆盖 P1 时的兜底文案"""
        report = dict(BASE_REPORT)
        report['coverage_matrix'] = {'matrix': {
            'IO': {'P1': {'covered': 9, 'total': 10}},
        }}
        prs = _open(generate_validation_summary_ppt(report))
        assert '各 Feature 覆盖率进展正常' in _all_text(prs)

    def test_charts_replaced(self):
        """PPT-UNIT-008: 传入图表截图后替换模板占位图"""
        prs = _open(generate_validation_summary_ppt(
            BASE_REPORT, chart_images={'feature': SAMPLE_PNG, 'p0p1': SAMPLE_PNG}))
        assert len(_pictures(prs)) == 2

    def test_charts_removed_when_missing(self):
        """PPT-UNIT-009: 无图表截图时移除占位图，不留样张假数据"""
        prs = _open(generate_validation_summary_ppt(BASE_REPORT, chart_images={}))
        assert len(_pictures(prs)) == 0

    def test_partial_charts(self):
        """PPT-UNIT-010: 只传一张图时只替换一张"""
        prs = _open(generate_validation_summary_ppt(
            BASE_REPORT, chart_images={'feature': SAMPLE_PNG}))
        assert len(_pictures(prs)) == 1

    def test_empty_project(self):
        """PPT-UNIT-011: 空项目（全 0 数据）不报错"""
        empty = {
            'project': {'name': 'Empty', 'start_date': '', 'end_date': ''},
            'overview': {'total_cp': 0, 'covered_cp': 0, 'coverage_rate': 0},
            'by_priority': {},
            'coverage_matrix': {'matrix': {}},
        }
        prs = _open(generate_validation_summary_ppt(empty))
        text = _all_text(prs)
        assert 'Empty-Functional Validation Summary' in text
        assert '0%' in text
        assert '0 / 0' in text

    def test_status_note_kept_as_placeholder(self):
        """PPT-UNIT-012: 右上状态文字保持模板占位文本（导出后用户编辑）"""
        prs = _open(generate_validation_summary_ppt(BASE_REPORT))
        assert '请在此填写项目当前状态与下一步计划' in _all_text(prs)
