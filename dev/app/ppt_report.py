"""Validation Summary PPT 生成器 (v0.16.0)

在导出项目材料包时，基于模板 validation_summary_template.pptx 自动生成
一页 "Functional Validation Summary" 幻灯片（版式与样张一致），放入材料包
ZIP 的 summary/ 子目录。

模板形状命名约定见 scripts/build_validation_summary_template.py。
本模块保持纯函数式：不访问数据库，数据由调用方（api.py）传入。
"""
import io
import os

from PIL import Image
from pptx import Presentation

TEMPLATE_PATH = os.path.join(
    os.path.dirname(__file__), "templates", "validation_summary_template.pptx"
)

# 评语规则：P1 覆盖率低于该阈值的 Feature 会被点名（与图表红色档位一致）
P1_LOW_RATE_THRESHOLD = 70
# 评语中最多列出的低覆盖 Feature 个数（避免文本溢出文本框）
COMMENTARY_MAX_FEATURES = 5


def _walk_shapes(shapes):
    """递归遍历形状（含组合形状内部）。"""
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:  # GROUP
            yield from _walk_shapes(shape.shapes)


def _set_text(shape, text):
    """设置文本并保留首个 run 的字体格式。"""
    tf = shape.text_frame
    para = tf.paragraphs[0]
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run._r.getparent().remove(run._r)
    else:
        para.text = text
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def _remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def _fmt_num(value):
    """数值格式化：92.6 -> '92.6'，100.0 -> '100'。"""
    return f"{round(float(value), 1):g}"


def _fmt_snapshot_date(date_str):
    """快照日期 '2026-05-28' -> '5/28'；无法解析时原样返回。"""
    try:
        parts = str(date_str).split("T")[0].split("-")
        return f"{int(parts[1])}/{int(parts[2])}"
    except (IndexError, ValueError):
        return str(date_str)


def _build_commentary(report_data):
    """按规则生成右下角评语。

    规则（对齐样张语义）：
    - P1 覆盖率 < 阈值的 Feature 点名 "的 P1 功能覆盖率不高"
    - P0 全部覆盖时补 "但 P0 核心功能已全部覆盖"，否则提示优先补齐
    """
    matrix = report_data.get("coverage_matrix", {}).get("matrix", {})

    low_p1 = []
    p0_incomplete = []
    has_p0 = False
    for feature, priorities in matrix.items():
        p1 = priorities.get("P1")
        if p1 and p1.get("total", 0) > 0:
            rate = round(p1.get("covered", 0) / p1["total"] * 100, 1)
            if rate < P1_LOW_RATE_THRESHOLD:
                low_p1.append((feature, rate))
        p0 = priorities.get("P0")
        if p0 and p0.get("total", 0) > 0:
            has_p0 = True
            if p0.get("covered", 0) < p0["total"]:
                p0_incomplete.append(feature)

    low_p1.sort(key=lambda item: item[1])

    parts = []
    if low_p1:
        names = "、".join(name for name, _ in low_p1[:COMMENTARY_MAX_FEATURES])
        parts.append(f"{names} 的 P1 功能覆盖率不高")
    if has_p0 and not p0_incomplete:
        parts.append("但 P0 核心功能已全部覆盖" if parts else "P0 核心功能已全部覆盖")
    elif has_p0:
        parts.append("P0 核心功能尚未全部覆盖，需优先补齐")

    return "；".join(parts) if parts else "各 Feature 覆盖率进展正常"


def _replace_chart(slide, shape, image_bytes):
    """用实际图表截图替换模板占位图，保持宽高比并居中于原占位区域。"""
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    _remove_shape(shape)

    if not image_bytes:
        return

    with Image.open(io.BytesIO(image_bytes)) as im:
        aspect = im.width / im.height
    box_aspect = width / height
    if aspect >= box_aspect:
        new_w = int(width)
        new_h = int(width / aspect)
    else:
        new_h = int(height)
        new_w = int(height * aspect)
    new_left = int(left + (width - new_w) / 2)
    new_top = int(top + (height - new_h) / 2)
    slide.shapes.add_picture(io.BytesIO(image_bytes), new_left, new_top, new_w, new_h)


def generate_validation_summary_ppt(report_data, coverage_delta=None, chart_images=None):
    """生成 Validation Summary PPT。

    参数:
        report_data: _get_project_report_data() 返回的数据 dict
        coverage_delta: None 或 {'prev_date': '2026-05-28', 'delta_pp': 0.8}，
                        为 None 时删除模板中的环比行
        chart_images: {'feature': png字节, 'p0p1': png字节}，缺失的图移除占位形状

    返回:
        pptx 文件字节
    """
    chart_images = chart_images or {}
    prs = Presentation(TEMPLATE_PATH)
    slide = prs.slides[0]
    shapes = {s.name: s for s in _walk_shapes(slide.shapes)}

    project = report_data.get("project", {})
    overview = report_data.get("overview", {})
    by_priority = report_data.get("by_priority", {})

    def set_text(name, text):
        shape = shapes.get(name)
        if shape is not None:
            _set_text(shape, text)

    # 标题: "{项目名}-Functional Validation Summary"
    set_text("title", f"{project.get('name', '')}-Functional Validation Summary")

    # KPI1: CP 覆盖率（简单占比口径）+ 环比
    set_text("kpi_coverage_value", f"{_fmt_num(overview.get('coverage_rate', 0))}%")
    delta_shape = shapes.get("kpi_coverage_delta")
    if delta_shape is not None:
        if coverage_delta:
            delta_pp = float(coverage_delta.get("delta_pp", 0))
            sign = "+" if delta_pp >= 0 else "-"
            text = f"较 {_fmt_snapshot_date(coverage_delta.get('prev_date', ''))} {sign}{_fmt_num(abs(delta_pp))}pp"
            _set_text(delta_shape, text)
        else:
            _remove_shape(delta_shape)

    # KPI2 / KPI3: 已覆盖 / 未覆盖 CP
    total_cp = int(overview.get("total_cp", 0))
    covered_cp = int(overview.get("covered_cp", 0))
    set_text("kpi_covered_value", str(covered_cp))
    set_text("kpi_uncovered_value", str(max(total_cp - covered_cp, 0)))

    # KPI4: P0 / P1 拆分（数量 + 占比，占比按 P0+P1 总数计算）
    p0_total = int(by_priority.get("P0", {}).get("total", 0))
    p1_total = int(by_priority.get("P1", {}).get("total", 0))
    p0p1_total = p0_total + p1_total
    p0_share = round(p0_total / p0p1_total * 100, 1) if p0p1_total > 0 else 0
    p1_share = round(p1_total / p0p1_total * 100, 1) if p0p1_total > 0 else 0
    set_text("kpi_p0p1_value", f"{p0_total} / {p1_total}")
    set_text("kpi_p0p1_split", f"P0={_fmt_num(p0_share)}%，P1={_fmt_num(p1_share)}%")

    # 右下评语（规则自动生成）；右上 status_note 为静态占位，不改动
    set_text("commentary", _build_commentary(report_data))

    # 两张图表：用材料包导出的图表截图替换模板占位图
    for shape_name, key in (("chart_feature", "feature"), ("chart_p0p1", "p0p1")):
        shape = shapes.get(shape_name)
        if shape is not None:
            _replace_chart(slide, shape, chart_images.get(key))

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()
